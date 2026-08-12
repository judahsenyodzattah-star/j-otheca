#!/usr/bin/env python3
"""Convert an academic Markdown doc to .docx (headings, bold/italic, bullets, numbered
lists, pipe tables, code blocks, hr). Usage: to_docx.py file.md [out.docx]"""
import re, sys
from docx import Document
from docx.shared import Pt

def add_rich(par, text):
    # **bold**, *italic*, `code` — simple inline markdown
    for i, tok in enumerate(re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith('`') and tok.endswith('`'):
            r = par.add_run(tok[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else:
            par.add_run(tok)

def convert(src, dst):
    doc = Document()
    lines = open(src, encoding='utf-8').read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run('\n'.join(code)); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
            continue
        m = re.match(r'^(#{1,6})\s+(.*)', ln)
        if m:
            doc.add_heading(re.sub(r'\*', '', m.group(2)).strip(), level=min(len(m.group(1)), 4))
            i += 1; continue
        if re.match(r'^\s*---+\s*$', ln):
            doc.add_paragraph('—' * 20); i += 1; continue
        if ln.lstrip().startswith('|'):  # pipe table
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.match(r'^:?-{3,}:?$', c.replace(' ', '')) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                t.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        t.rows[ri].cells[ci].text = re.sub(r'\*\*', '', cell)
            continue
        m = re.match(r'^\s*[-*]\s+(.*)', ln)
        if m:
            add_rich(doc.add_paragraph(style='List Bullet'), m.group(1))
            i += 1; continue
        m = re.match(r'^\s*\d+[.)]\s+(.*)', ln)
        if m:
            add_rich(doc.add_paragraph(style='List Number'), m.group(1))
            i += 1; continue
        p = doc.add_paragraph()
        add_rich(p, ln)
        i += 1
    doc.save(dst)
    print('wrote', dst)

if __name__ == '__main__':
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else re.sub(r'\.md$', '.docx', src)
    convert(src, dst)
